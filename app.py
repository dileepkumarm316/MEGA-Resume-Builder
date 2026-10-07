import streamlit as st
from docx import Document
import io

st.title("🚀 AI Resume Maker PRO")
st.write("By Dileep Kumar M - PRO")

name = st.text_input("Name", "Dileep")
role = st.text_input("Role", "Software Engineer")
skills = st.text_area("Skills", "Python, SQL")
exp = st.text_area("Exp", "B.Tech CSE")

if st.button("Generate PRO Resume"):
    st.success("PRO Resume Ready!")
    st.write(name, "-", role)
    st.write(skills)
    st.write(exp)
    doc = Document()
    doc.add_heading(name, 0)
    doc.add_paragraph(role)
    doc.add_paragraph(skills)
    doc.add_paragraph(exp)
    bio = io.BytesIO()
    doc.save(bio)
    st.download_button("Download",bio.getvalue(),"a.docx")